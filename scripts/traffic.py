#!/usr/bin/env python3
r"""traffic.py — keep Cloudflare's request log for the sites in a local database, and say who is in it.

*(Bill, 2026-10-29.)* Cloudflare's Free plan keeps 31 days of per-request records, readable
through the GraphQL `httpRequestsAdaptive` dataset with a read-only token in `CF_API_TOKEN`.
`pull` copies every complete UTC day not yet held, for both zones, into a SQLite file outside
the repository — `DB`, overridable by `CORPUS_TRAFFIC_DB` — because it holds visitors' IP
addresses, which are not the repository's to publish. The poll runs it each tick; a day
already held costs no API call, so the night's first tick does the work and the rest are free.

**People are kept for good; machines for `RAW_DAYS`.** Who is a person can only be decided
from the whole day — the beacon, the 404 count — so every request is pulled. After `RAW_DAYS`
the automated rows are deleted and survive as one count per class and source a day in
`tallies`; the rows left classify as they did, because what made them person or unknown is
still there (Bill, 2026-10-29: "we only really need real users"). Raw, a month is ~80 MB.

Only `requestSource: eyeball` is kept. The site is served by a Worker, and every request it
answers from static assets is logged a second time as `edgeWorkerFetch` — the same ray,
nothing the visitor did. Counting both doubles the traffic.

`report` classes each request and tallies a day. The classes are computed when read, not
stored, so a rule can be corrected and the whole history re-read under it:

- **bot** — Cloudflare names it a verified bot, or the user agent says it is one;
- **scanner** — an IP with `SCANNER_404S` or more 404s in the day: probes for `.env`, `.git`;
- **cloud** — a data-centre network (`CLOUD_ASN`) wearing a browser's user agent, which is
  how undeclared scrapers and headless crawlers arrive;
- **person** — anything else that ran the page's JavaScript, i.e. sent the `/cdn-cgi/rum`
  analytics beacon that day; the one sign a browser was really rendering the page;
- **unknown** — the rest: a residential address that never sent the beacon (a person with
  a blocker, or a scraper on a residential proxy);
- **own** — Bill's own networks (`OWN_FILE`), checked first, so his views and the site
  checks run from this machine are never readers.

**`logs/site-analytics.xlsx` is people only** (Bill, 2026-10-29): *Access*, a row a day of
users (distinct IPs) and pages viewed; *Pages*, a row a page with its views a month. A page
is a successful GET of HTML — not an asset, a download or the beacon. `pull` rewrites it
whenever the database is newer; open in Excel it cannot be replaced, so it waits for the
next tick. It replaced Google Analytics' `site-analytics.csv`, and like it is committed by the
cycle (`CYCLE.md` 1a) — the one copy of the readers' history that git holds.

Usage:  python scripts/traffic.py pull              # every complete day not yet held
        python scripts/traffic.py pull --quiet      # print only when a day was added
        python scripts/traffic.py report [YYYY-MM-DD]   # default: the latest day held
        python scripts/traffic.py xlsx              # rewrite the workbook now
"""
from __future__ import annotations

import argparse
import collections
import datetime as dt
import json
import os
import re
import sqlite3
import sys
import urllib.error
import urllib.request

DB = os.environ.get("CORPUS_TRAFFIC_DB", r"C:\corpus-traffic\traffic.sqlite")
XLSX = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                    "logs", "site-analytics.xlsx")
API = "https://api.cloudflare.com/client/v4/graphql"
ZONES = {"data-landscapers.io": "e6663237e18dfc5241e44003a2b1d976",
         "data-landscapers.com": "dacc4524150b875a9d9aca87909eb02f"}
KEEP_DAYS = 30          # the dataset's notOlderThan is 31 days; the oldest is a partial day
RAW_DAYS = 90           # automated rows older than this become tallies
PAGE = 10000            # the dataset's maxPageSize; a window that fills it is split in two

FIELDS = ("rayName datetime clientIP clientAsn clientASNDescription clientCountryName "
          "clientRequestHTTPHost clientRequestHTTPMethodName clientRequestPath clientRequestQuery "
          "edgeResponseStatus cacheStatus userAgent verifiedBotCategory clientDeviceType").split()
QUERY = ("query($z:String!,$a:Time!,$b:Time!){viewer{zones(filter:{zoneTag:$z}){"
         f"httpRequestsAdaptive(limit:{PAGE},filter:{{datetime_geq:$a,datetime_lt:$b,"
         "requestSource:\"eyeball\"},orderBy:[datetime_ASC]){" + " ".join(FIELDS) + "}}}}")

BOT_UA = re.compile(r"bot|crawl|spider|slurp|fetch|scan|preview|python|curl|wget|go-http|"
                    r"java/|okhttp|axios|headless|httpclient|libwww|monitor|uptime", re.I)
CLOUD_ASN = re.compile(r"google|amazon|microsoft|tencent|alibaba|huawei|leaseweb|constant|"
                       r"latitude|hostroyale|digitalocean|ovh|hetzner|linode|akamai|oracle|"
                       r"cloudflare|contabo|vultr|m247|datacamp|hostinger|scaleway", re.I)
SCANNER_404S = 20
CLASSES = ("person", "unknown", "bot", "cloud", "scanner", "own")
# Bill's own networks, as address prefixes one a line, kept beside the database rather than in
# the repository because they locate his home. His browser blocks the beacon, so he would
# read as unknown anyway; this makes it a rule rather than an accident (Bill, 2026-10-29).
OWN_FILE = os.path.join(os.path.dirname(DB), "own-networks.txt")


def own_networks() -> tuple[str, ...]:
    try:
        with open(OWN_FILE, encoding="utf-8") as fh:
            return tuple(l.split("#")[0].strip() for l in fh if l.split("#")[0].strip())
    except OSError:
        return ()


def token() -> str:
    """The environment, or the registry when the poll's tab predates the variable."""
    t = os.environ.get("CF_API_TOKEN")
    if not t and sys.platform == "win32":
        import winreg
        try:
            with winreg.OpenKey(winreg.HKEY_CURRENT_USER, "Environment") as k:
                t = winreg.QueryValueEx(k, "CF_API_TOKEN")[0]
        except OSError:
            pass
    if not t:
        sys.exit("traffic: CF_API_TOKEN is not set")
    return t


def fetch(zone: str, a: dt.datetime, b: dt.datetime, tok: str) -> list[dict]:
    body = json.dumps({"query": QUERY, "variables": {
        "z": zone, "a": a.strftime("%Y-%m-%dT%H:%M:%SZ"), "b": b.strftime("%Y-%m-%dT%H:%M:%SZ")}})
    req = urllib.request.Request(API, data=body.encode(), headers={
        "Authorization": f"Bearer {tok}", "Content-Type": "application/json"})
    try:
        with urllib.request.urlopen(req, timeout=120) as r:
            d = json.load(r)
    except urllib.error.HTTPError as e:
        d = json.load(e)
    if d.get("errors"):
        raise RuntimeError(d["errors"][0].get("message"))
    rows = d["data"]["viewer"]["zones"][0]["httpRequestsAdaptive"]
    if len(rows) >= PAGE:
        mid = a + (b - a) / 2
        return fetch(zone, a, mid, tok) + fetch(zone, mid, b, tok)
    return rows


def connect() -> sqlite3.Connection:
    os.makedirs(os.path.dirname(DB), exist_ok=True)
    con = sqlite3.connect(DB)
    con.execute("CREATE TABLE IF NOT EXISTS requests (zone TEXT, day TEXT, "
                + ", ".join(f"{f} TEXT" for f in FIELDS) + ", PRIMARY KEY (rayName))")
    con.execute("CREATE INDEX IF NOT EXISTS requests_day ON requests (day)")
    con.execute("CREATE TABLE IF NOT EXISTS days (zone TEXT, day TEXT, rows INTEGER, "
                "pulled TEXT, PRIMARY KEY (zone, day))")
    con.execute("CREATE TABLE IF NOT EXISTS tallies (day TEXT, class TEXT, source TEXT, "
                "requests INTEGER, ips INTEGER, PRIMARY KEY (day, class, source))")
    return con


AUTOMATED = ("bot", "cloud", "scanner")


def source(r: sqlite3.Row, k: str) -> str:
    """The name a tally is kept under: the bot's own name where it gives one, else its network."""
    ua = re.search(r"([A-Za-z-]*(?:bot|crawler|spider|user)[A-Za-z-]*)/?", r["userAgent"], re.I)
    return ua.group(1) if ua and k == "bot" else r["clientASNDescription"][:40]


def prune(con: sqlite3.Connection) -> int:
    """Turn the automated rows of each day older than RAW_DAYS into tallies."""
    con.row_factory = sqlite3.Row
    cut = (dt.datetime.now(dt.timezone.utc).date() - dt.timedelta(days=RAW_DAYS)).isoformat()
    done = {d for (d,) in con.execute("SELECT DISTINCT day FROM tallies")}
    days = [d for (d,) in con.execute("SELECT DISTINCT day FROM requests WHERE day < ?", (cut,))
            if d not in done]
    for day in days:
        rows = con.execute("SELECT * FROM requests WHERE day=?", (day,)).fetchall()
        groups = collections.defaultdict(set)
        n = collections.Counter()
        for r, k in zip(rows, classify(rows)):
            if k in AUTOMATED:
                groups[(k, source(r, k))].add(r["clientIP"])
                n[(k, source(r, k))] += 1
                con.execute("DELETE FROM requests WHERE rayName=?", (r["rayName"],))
        con.executemany("INSERT INTO tallies VALUES (?,?,?,?,?)",
                        [(day, k, s, n[(k, s)], len(ips)) for (k, s), ips in groups.items()])
        con.commit()
    con.row_factory = None
    return len(days)


def pull(quiet: bool) -> int:
    con, tok = connect(), None
    today = dt.datetime.now(dt.timezone.utc).date()
    added = []
    for name, zone in ZONES.items():
        held = {d for (d,) in con.execute("SELECT day FROM days WHERE zone=?", (name,))}
        for n in range(KEEP_DAYS, 0, -1):
            day = today - dt.timedelta(days=n)
            if day.isoformat() in held:
                continue
            tok = tok or token()
            a = dt.datetime.combine(day, dt.time(), dt.timezone.utc)
            rows = []
            for h in range(0, 24, 6):
                rows += fetch(zone, a + dt.timedelta(hours=h), a + dt.timedelta(hours=h + 6), tok)
            con.executemany(
                f"INSERT OR IGNORE INTO requests VALUES ({','.join('?' * (len(FIELDS) + 2))})",
                [(name, day.isoformat(), *(str(r[f]) for f in FIELDS)) for r in rows])
            con.execute("INSERT INTO days VALUES (?,?,?,?)", (name, day.isoformat(), len(rows),
                        dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds")))
            con.commit()
            added.append(f"{name} {day} {len(rows)}")
    if prune(con):
        con.execute("VACUUM")
    con.close()
    note = ""
    if not os.path.exists(XLSX) or os.path.getmtime(XLSX) < os.path.getmtime(DB):
        try:
            workbook()
        except PermissionError:
            note = " · site-analytics.xlsx is open in Excel, rewritten next tick"
    if added or note or not quiet:
        print(f"traffic: {len(added)} day(s) added" + (" — " + "; ".join(added[-3:]) if added else "") + note)
    return 0


PAGE_PATH = re.compile(r"(/|\.html?|/[^./]+)$")


def page(r: sqlite3.Row) -> str | None:
    """The page a request viewed, as host and path, or None when it was not a page view."""
    p = r["clientRequestPath"]
    if (r["clientRequestHTTPMethodName"] != "GET" or r["edgeResponseStatus"] not in ("200", "304")
            or p.startswith(("/assets/", "/cdn-cgi/")) or not PAGE_PATH.search(p)):
        return None
    host = r["clientRequestHTTPHost"].removeprefix("www.")
    return host + re.sub(r"index\.html?$", "", p)


def workbook() -> None:
    """Write XLSX from the person rows: Access a row a day, Pages a row a page by month."""
    import openpyxl
    from openpyxl.styles import Font
    con = connect()
    con.row_factory = sqlite3.Row
    users, views = collections.defaultdict(set), collections.Counter()
    months, by_page = set(), collections.defaultdict(collections.Counter)
    days = sorted({d for (d,) in con.execute("SELECT DISTINCT day FROM days")})
    for day in days:
        rows = con.execute("SELECT * FROM requests WHERE day=?", (day,)).fetchall()
        for r, k in zip(rows, classify(rows)):
            if k != "person":
                continue
            users[day].add(r["clientIP"])
            if (pg := page(r)):
                views[day] += 1
                months.add(day[:7])
                by_page[pg][day[:7]] += 1
    con.close()
    months = sorted(months)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Access"
    ws.append(["Date", "Users", "Pages"])
    for day in days:
        ws.append([dt.date.fromisoformat(day), len(users[day]), views[day]])
        ws.cell(ws.max_row, 1).number_format = "yyyy-mm-dd"
    ws.column_dimensions["A"].width = 12
    ps = wb.create_sheet("Pages")
    ps.append(["Page"] + [dt.date.fromisoformat(m + "-01").strftime("%b %Y") for m in months])
    latest = months[-1] if months else ""
    for pg, n in sorted(by_page.items(), key=lambda x: (-x[1][latest], -sum(x[1].values()), x[0])):
        ps.append([pg] + [n[m] for m in months])
    ps.column_dimensions["A"].width = 70
    for sheet in (ws, ps):
        sheet.freeze_panes = "A2"
        for c in sheet[1]:
            c.font = Font(bold=True)
    tmp = XLSX + ".tmp"
    wb.save(tmp)
    try:
        os.replace(tmp, XLSX)
    except PermissionError:
        os.remove(tmp)
        raise


def classify(rows: list[sqlite3.Row]) -> list[str]:
    fours = collections.Counter(r["clientIP"] for r in rows if r["edgeResponseStatus"] == "404")
    beacon = {r["clientIP"] for r in rows if r["clientRequestPath"] == "/cdn-cgi/rum"}
    own = own_networks()
    out = []
    for r in rows:
        if own and r["clientIP"].startswith(own):
            out.append("own")
        elif r["verifiedBotCategory"] or BOT_UA.search(r["userAgent"]):
            out.append("bot")
        elif fours[r["clientIP"]] >= SCANNER_404S:
            out.append("scanner")
        elif CLOUD_ASN.search(r["clientASNDescription"]):
            out.append("cloud")
        elif r["clientIP"] in beacon:
            out.append("person")
        else:
            out.append("unknown")
    return out


def report(day: str | None) -> int:
    con = connect()
    con.row_factory = sqlite3.Row
    day = day or con.execute("SELECT max(day) FROM days").fetchone()[0]
    if not day:
        sys.exit("traffic: nothing held yet; run pull")
    rows = con.execute("SELECT * FROM requests WHERE day=?", (day,)).fetchall()
    cls = classify(rows)
    tally = collections.Counter(cls)
    ips = {c: len({r["clientIP"] for r, k in zip(rows, cls) if k == c}) for c in CLASSES}
    who = collections.Counter()
    for r, k in zip(rows, cls):
        if k in AUTOMATED:
            who[(k, source(r, k))] += 1
    for t in con.execute("SELECT * FROM tallies WHERE day=?", (day,)):
        who[(t["class"], t["source"])] += t["requests"]
        tally[t["class"]] += t["requests"]
        ips[t["class"]] += t["ips"]   # a pruned day's IPs are summed by source, so may overcount
    print(f"{day}: {sum(tally.values())} requests — " + " ; ".join(
        f"{c} {tally[c]} from {ips[c]} IPs" for c in CLASSES))
    print("\nbiggest automated sources:")
    for (k, name), n in who.most_common(12):
        print(f"  {n:6}  {k:8} {name}")
    pages = collections.Counter(r["clientRequestPath"] for r, k in zip(rows, cls)
                                if k == "person" and not r["clientRequestPath"].startswith(("/assets/", "/cdn-cgi/")))
    print("\nwhat people opened:")
    for p, n in pages.most_common(15):
        print(f"  {n:6}  {p}")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description="Cloudflare request log: pull to SQLite, report a day.")
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("pull")
    p.add_argument("--quiet", action="store_true", help="print only when a day was added")
    r = sub.add_parser("report")
    r.add_argument("day", nargs="?")
    sub.add_parser("xlsx")
    args = ap.parse_args()
    if args.cmd == "xlsx":
        workbook()
        return 0
    return pull(args.quiet) if args.cmd == "pull" else report(args.day)


if __name__ == "__main__":
    sys.exit(main())
